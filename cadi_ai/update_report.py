import docx
import os

docx_path = "CADI_AI_Report (1).docx"
doc = docx.Document(docx_path)

# Update Table 3 caption
for p in doc.paragraphs:
    if "Indicates: CBAM attention provides consistent broad-spectrum improvement at negligible parameter cost." in p.text:
        p.text = p.text.replace("Indicates: CBAM attention provides consistent broad-spectrum improvement at negligible parameter cost.", "Indicates: CBAM attention provides consistent broad-spectrum improvement at negligible parameter cost. The precision decrease (−0.010) reflects increased model sensitivity — CBAM detects more true positives but also accepts slightly more false positives, a deliberate and operationally appropriate trade-off in agricultural disease monitoring where missed infections (false negatives) carry a higher cost than false alarms (false positives).")

# Add references
doc.add_paragraph("[6] Jocher, G., Chaurasia, A., & Qiu, J. (2023). Ultralytics YOLO (Version 8.0.0) [Software].\n    GitHub. https://github.com/ultralytics/ultralytics")
doc.add_paragraph("[7] Lin, T. Y., Goyal, P., Girshick, R., He, K., & Dollár, P. (2017).\n    Focal Loss for Dense Object Detection. ICCV 2017. arXiv:1708.02002.")

doc.save("CADI_AI_Report (1).docx")
print("Report updated.")